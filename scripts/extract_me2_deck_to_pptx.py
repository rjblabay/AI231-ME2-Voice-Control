"""Create an editable PPTX from the supplied two-page PDF template.

The original rendered pages are retained as slide backgrounds; the filled
fields are real editable PowerPoint shapes placed over the PDF placeholders.
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
ASSETS = Path('/private/tmp/me2-pdf-render')
OUT = ROOT / 'MEX2_Slides' / 'me2-deck-editable.pptx'
W, H = Inches(841.889764 / 72), Inches(595.275591 / 72)

NAVY = '172C44'; LIGHT = 'F3F5F9'; BLACK = '000000'; TEXT = '15243A'
PALE_BLUE = 'F2F5FB'; PALE_GREEN = 'E5F0E6'; PALE_YELLOW = 'FFF2D8'; PALE_PURPLE = 'EEE9FC'

def rgb(hex_value):
    return RGBColor.from_string(hex_value)

def box(slide, x, y, w, h, value, fill, colour=TEXT, size=8.5, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x/72), Inches(y/72), Inches(w/72), Inches(h/72))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    frame = shape.text_frame; frame.clear(); frame.margin_left = Pt(3); frame.margin_right = Pt(2); frame.margin_top = Pt(1); frame.margin_bottom = Pt(0)
    para = frame.paragraphs[0]; para.text = value; para.alignment = align
    run = para.runs[0]; run.font.name = 'Arial'; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = rgb(colour)
    return shape

def page_one(slide):
    # Model pills: fill only their inner text areas.
    box(slide, 35, 155, 372, 22, 'Mic · 16 kHz → 40-bin log-Mel × 300 frames', PALE_BLUE, size=8.4, bold=True)
    box(slide, 35, 186, 372, 21, 'CRNN encoder · 3 CNN blocks + 2-layer BiGRU', PALE_GREEN, '244827', 8.4, True)
    box(slide, 95, 209, 245, 10, 'streaming inference', BLACK, '397524', 6.8, False, PP_ALIGN.CENTER)
    box(slide, 36, 220, 176, 20, 'Intent classifier', PALE_YELLOW, '674010', 8.0, True, PP_ALIGN.CENTER)
    box(slide, 227, 220, 177, 20, 'Phrase / slot classifier', PALE_YELLOW, '674010', 7.7, True, PP_ALIGN.CENTER)
    box(slide, 36, 251, 176, 20, 'Local action handler', PALE_PURPLE, '3A2B78', 8.0, True, PP_ALIGN.CENTER)

    dataset = ['AI231 Hugging Face + custom Astra', '16,868 total clips', 'mixed speakers + synthetic voices', '22 intents · 99 output classes']
    for y, value in zip((181, 210, 239, 266), dataset): box(slide, 556, y, 248, 18, value, LIGHT, size=7.6)
    training = ['UP HPC GPU / CUDA', 'weighted CE (cmd) + focal loss (wake)', 'Adam + cosine LR schedule', '66 command / 80 wake epochs']
    for y, value in zip((354, 383, 412, 441), training): box(slide, 555, y, 248, 18, value, LIGHT, size=7.2)

    box(slide, 228, 315, 180, 18, 'Cmd 1.35M / 5.2MB · Wake 0.65M / 2.5MB', LIGHT, size=6.4)
    box(slide, 229, 407, 179, 18, '98.44% / 90.16%', LIGHT, size=8.3)
    # Change this one existing label to a more meaningful metric while retaining the row.
    box(slide, 30, 435, 190, 18, 'Wake / intent macro-F1', BLACK, 'A9C2DF', size=8.5)
    box(slide, 229, 435, 179, 18, '82.93% / 89.20%', LIGHT, size=8.3)
    box(slide, 229, 465, 179, 18, 'Not yet measured', LIGHT, size=8.0)
    box(slide, 229, 494, 179, 18, 'ONNX Runtime · RPi 4', LIGHT, size=8.0)

def page_two(slide):
    values = ['Pending: publish project repository', 'AI231 HF dataset + custom Astra data', 'UP HPC GPU / CUDA · 66 / 80 epochs', 'Saved locally: models/macbook/*.onnx']
    for y, value in zip((83, 128, 174, 219), values): box(slide, 226, y, 580, 30, value, BLACK, 'F5F8FF', 8.5)
    statuses = ['PENDING', 'PENDING', 'LOCAL ONLY', 'PENDING', 'COMPLETE', 'NOT INCLUDED']
    for y, value in zip((322, 353, 384, 415, 446, 477), statuses): box(slide, 695, y, 82, 18, value, LIGHT, size=7.2, bold=True, align=PP_ALIGN.CENTER)

def main():
    deck = Presentation(); deck.slide_width = W; deck.slide_height = H
    blank = deck.slide_layouts[6]
    for index, image in enumerate((ASSETS/'original-1.png', ASSETS/'original-2.png')):
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=W, height=H)
        (page_one if index == 0 else page_two)(slide)
    deck.save(OUT)
    print(OUT)

if __name__ == '__main__': main()
