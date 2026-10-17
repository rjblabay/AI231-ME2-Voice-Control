"""Create a clean editable PPTX container from the original ME2 PDF pages.

The PDF render is retained as a locked-looking background image on each slide;
new content will be added as editable shapes only after the content map is
approved.
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
RENDER_DIR = Path("/private/tmp/me2-pdf-render")
OUT = ROOT / "MEX2_Slides" / "me2-deck-editable-template.pptx"


def main() -> None:
    deck = Presentation()
    deck.slide_width = Inches(841.889764 / 72)
    deck.slide_height = Inches(595.275591 / 72)
    blank = deck.slide_layouts[6]
    for page in ("original-1.png", "original-2.png"):
        slide = deck.slides.add_slide(blank)
        slide.shapes.add_picture(str(RENDER_DIR / page), 0, 0,
                                 width=deck.slide_width, height=deck.slide_height)
    deck.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
